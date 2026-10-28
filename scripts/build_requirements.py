"""Собрать Word с техническими требованиями из docs/requirements/REQUIREMENTS.md.

Из корня проекта: .venv/bin/python -m scripts.build_requirements
Зависимости: requirements-docs.txt. PDF экспортируется из Word или LibreOffice.
"""

from datetime import datetime, timezone
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
DIRECTORY = ROOT / "docs" / "requirements"
SOURCE = DIRECTORY / "REQUIREMENTS.md"
OUTPUT = DIRECTORY / "technical_requirements.docx"


def set_font(style, size, *, bold=False):
    style.font.name = "Times New Roman"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().rFonts
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), "Times New Roman")
    for name in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn(f"w:{name}"), None)


def add_table(document, lines):
    rows = [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]
    headers = rows[0]
    if len(headers) != 3 or len(rows) < 3:
        raise ValueError("Ожидается таблица с тремя столбцами, заголовком и строками данных")
    if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in rows[1]):
        raise ValueError("После заголовка таблицы должна идти строка разделителей")
    data = [headers, *rows[2:]]
    if any(len(row) != len(headers) for row in data):
        raise ValueError("Число ячеек должно совпадать с числом столбцов")

    section = document.sections[0]
    available_mm = (section.page_width - section.left_margin - section.right_margin) / 36000
    if headers[0] == "Атрибут":
        first, second = 42, 35
    elif headers[0] == "Код и атрибут":
        first, second = 46, 57
    else:
        first, second = 20, 40
    widths = (first, second, available_mm - first - second)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)

    for row_index, values in enumerate(data):
        row = table.rows[0] if row_index == 0 else table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for cell_index, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_with_next = row_index == 0
            run = paragraph.add_run(value)
            run.font.name, run.font.size = "Times New Roman", Pt(12)
            run.bold = row_index == 0 or cell_index == 0
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def build():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(25.4)
    section.left_margin = section.right_margin = Mm(25.4)

    normal = document.styles["Normal"]
    set_font(normal, 14)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.widow_control = True
    for name, size in (("Title", 16), ("Heading 1", 14), ("Heading 2", 13)):
        style = document.styles[name]
        set_font(style, size, bold=True)
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
    document.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        index += 1
        if not line.strip():
            continue
        if line.startswith("|"):
            table_lines = [line]
            while index < len(lines) and lines[index].startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(document, table_lines)
        elif line.startswith("# "):
            document.add_paragraph(line[2:], "Title")
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            paragraph = document.add_heading(line[4:], level=2)
            paragraph.paragraph_format.left_indent = Mm(6.35)
        else:
            paragraph = document.add_paragraph()
            match = re.match(r"^((?:ПТ-\d+\.\d+|БТ-\d+|ФТ-\d+|БП-\d+|НФТ-\d+)\.) (.+)$", line)
            if match:
                paragraph.add_run(match[1] + " ").bold = True
                paragraph.add_run(match[2])
                if match[1].startswith("ПТ-"):
                    paragraph.paragraph_format.left_indent = Mm(6.35)
                paragraph.paragraph_format.keep_together = True
            else:
                paragraph.add_run(line)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, value, end):
        run._r.append(node)
    run.font.name, run.font.size = "Times New Roman", Pt(12)

    document.core_properties.title = "Технические требования к модулю управления каталогом товаров"
    document.core_properties.subject = "Курсовой проект. Пользовательские, функциональные и нефункциональные требования"
    document.core_properties.author = "Чемоданов Никита Сергеевич"
    document.core_properties.last_modified_by = "Сборка проекта"
    document.core_properties.created = document.core_properties.modified = datetime.now(timezone.utc)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
