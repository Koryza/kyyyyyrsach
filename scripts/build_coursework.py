"""Собрать редактируемые DOCX и PPTX из материалов проекта.

Необязательные зависимости: requirements-docs.txt. Из корня: python -m scripts.build_coursework
PDF получается отдельным экспортом LibreOffice/Word/PowerPoint; это не часть сервера.
"""

import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt as WordPt, RGBColor as WordColor
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
WORD_FILE = DOCS / "coursework" / "coursework_draft.docx"
SLIDES_FILE = DOCS / "defense" / "catalog_defense.pptx"
BLUE, INK, GRAY, BACK, WHITE = "345DE3", "18253B", "64748B", "F5F7FB", "FFFFFF"
STUDENT = json.loads((DOCS / "coursework" / "student.json").read_text())


def plain(text):
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.replace("**", "").replace("`", "")


def field(paragraph, instruction, cached=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText"); code.set(qn("xml:space"), "preserve"); code.text = instruction
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t"); value.text = cached
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for element in (begin, code, separate, value, end): run._r.append(element)


def make_word():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin, section.right_margin = Mm(30), Mm(10)
    section.top_margin = section.bottom_margin = Mm(20)
    section.header_distance = Mm(10)
    section.different_first_page_header_footer = True
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", WordPt(14)
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        normal._element.rPr.rFonts.attrib.pop(qn("w:" + attribute), None)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Mm(12.5)
    normal.paragraph_format.space_after = WordPt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Heading 1", "Heading 2"):
        style = document.styles[name]
        style.font.name, style.font.size, style.font.bold = "Times New Roman", WordPt(14), True
        style.font.color.rgb = WordColor(0, 0, 0)
        for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            style._element.rPr.rFonts.attrib.pop(qn("w:" + attribute), None)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        style.paragraph_format.first_line_indent = Mm(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = WordPt(18)
        style.paragraph_format.space_after = WordPt(14)
        style.paragraph_format.line_spacing = 1.5
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.paragraph_format.first_line_indent = Mm(0)
    field(header, " PAGE ")

    def cover(text, *, bold=False, before=0, after=0, size=14, align=WD_ALIGN_PARAGRAPH.CENTER):
        paragraph = document.add_paragraph()
        paragraph.alignment = align
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.space_before = WordPt(before)
        paragraph.paragraph_format.space_after = WordPt(after)
        run = paragraph.add_run(text); run.bold = bold; run.font.size = WordPt(size)

    cover("Аккредитованное образовательное частное учреждение\nвысшего образования")
    cover("«Московский финансово-юридический университет МФЮА»\n(МФЮА)", bold=True)
    cover(f"КАФЕДРА «{STUDENT['department'].upper()}»", before=14)
    cover("КУРСОВАЯ РАБОТА", bold=True, before=35, after=12)
    cover(f"по дисциплине: «{STUDENT['discipline']}»", after=12)
    cover("на тему: «Разработка модуля управления каталогом\nтоваров для интернет-магазина»", bold=True, after=25)
    cover(f"Специальность: {STUDENT['specialty_code']}\n{STUDENT['specialty']}\nПрофиль: {STUDENT['profile']}", align=WD_ALIGN_PARAGRAPH.LEFT, after=18)
    cover(f"Автор работы: {STUDENT['student_name']}\nКурс: {STUDENT['course']}       Группа: {STUDENT['group']} ({STUDENT['group_short']})\nРуководитель работы: {STUDENT['supervisor']}\nПодписи и даты: ________________________", align=WD_ALIGN_PARAGRAPH.LEFT, after=25)
    cover(f"{STUDENT['city']} — {STUDENT['year']}", before=25)
    document.add_page_break()
    document.add_heading("СОДЕРЖАНИЕ", 1)
    toc_anchor = document.add_paragraph()
    headings = []
    document.add_page_break()

    def heading(title, level=1):
        paragraph = document.add_heading(title.upper(), level)
        if level == 1 and headings: paragraph.paragraph_format.page_break_before = True
        bookmark = f"chapter_{len(headings) + 1}"
        start = OxmlElement("w:bookmarkStart"); start.set(qn("w:id"), str(100 + len(headings))); start.set(qn("w:name"), bookmark)
        end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"), str(100 + len(headings)))
        paragraph._p.insert(0, start); paragraph._p.append(end)
        headings.append((title, bookmark, level))

    def table(caption, headers, rows, widths):
        paragraph = document.add_paragraph(caption)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.paragraph_format.space_before = WordPt(12)
        paragraph.paragraph_format.keep_with_next = True
        grid = document.add_table(rows=1, cols=len(headers))
        grid.style = "Table Grid"
        grid.autofit = False
        for column, width in zip(grid.columns, widths): column.width = Mm(width)
        repeating = OxmlElement("w:tblHeader")
        grid.rows[0]._tr.get_or_add_trPr().append(repeating)
        for cell, value in zip(grid.rows[0].cells, headers): cell.text = value
        for row in rows:
            for cell, value in zip(grid.add_row().cells, row): cell.text = value
        for row_index, row in enumerate(grid.rows):
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell, width in zip(row.cells, widths):
                cell.width = Mm(width)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.first_line_indent = Mm(0)
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    paragraph.paragraph_format.keep_with_next = row_index < len(grid.rows) - 1
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = WordPt(14)
                        run.bold = row_index == 0
        document.add_paragraph()

    def diagram(kind):
        if kind == "model":
            values = ["КАТЕГОРИЯ\nid (PK)\nname\nname_normalized\ndescription", "1 : N\n→", "ТОВАР\nid (PK) · sku (UQ)\ncategory_id (FK)\nname · description\nprice_kopecks · quantity\ncreated_at · updated_at"]
            widths = (57, 21, 92)
            caption = "Рисунок 1 — Логическая структура данных каталога"
        else:
            values = ["Браузер\nФормы", "→", "FastAPI\nHTML / API", "→", "catalog.py\nОперации", "→", "SQLAlchemy\nSession", "→", "SQLite\nДанные"]
            widths = (28, 7.5, 28, 7.5, 28, 7.5, 28, 7.5, 28)
            caption = "Рисунок 2 — Обработка запроса в приложении"
        drawing = document.add_table(rows=1, cols=len(values))
        drawing.autofit = False
        for column, width in zip(drawing.columns, widths): column.width = Mm(width)
        for index, (cell, value, width) in enumerate(zip(drawing.rows[0].cells, values, widths)):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text = value
            if index % 2 == 0:
                properties = cell._tc.get_or_add_tcPr()
                borders = OxmlElement("w:tcBorders")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:val"), "single"); edge.set(qn("w:sz"), "6")
                    borders.append(edge)
                properties.append(borders)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Mm(0)
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.line_spacing = 1.15
                for run in paragraph.runs:
                    run.font.name = "Times New Roman"; run.font.size = WordPt(12)
        paragraph = document.add_paragraph(caption)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.paragraph_format.keep_with_next = True
        paragraph = document.add_paragraph("Источник: составлено по модели и исходному коду проекта.")
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    lines = (DOCS / "coursework" / "DRAFT.md").read_text().splitlines()
    active = False
    for line in lines:
        if line.startswith("## Введение"): active = True
        if line.startswith("## Что осталось"): break
        if not active or not line.strip(): continue
        if line.startswith("<!-- figure:"):
            diagram(line.split(":", 1)[1].split()[0])
        elif line.startswith("## "):
            heading(line[3:].replace(" — рабочая основа", ""))
        elif line.startswith("### "):
            heading(line[4:], 2)
        elif not line.startswith(("Этот фрагмент введения", "Иллюстрации для этого параграфа:")):
            document.add_paragraph(plain(line))
            if line.startswith("Отдельный веб-модуль"):
                table("Таблица 1 — Проектное сравнение подходов", ("Подход", "Что организуется", "Соответствие задаче"), [
                    ("Таблица товаров", "Столбцы свойств, правила ввода и связи нужно настроить", "Подходит для исходного перечня; HTTP-модуль не реализует"),
                    ("Готовая система", "Выбор продукта, настройка или расширение", "Возможный прикладной путь; самостоятельная разработка не показана"),
                    ("Отдельный модуль", "Собственные схема, операции, формы и проверки", "Выбран для изучения Python и FastAPI"),
                ], (35, 65, 70))
            elif line.startswith("Поиск сочетается"):
                table("Таблица 2 — Связь требований с проверками", ("Требование", "Реализация", "Признак корректности"), [
                    ("CRUD товара", "Форма, карточка, API", "Изменения видны после чтения; ID сохраняется"),
                    ("Категории", "Справочник и связь", "Повторное имя и удаление занятой категории отклоняются"),
                    ("Поиск", "Нормализованное имя", "Русский регистр не меняет результат"),
                    ("Фильтры", "Условия общего запроса", "Категория, цена и наличие применяются вместе"),
                    ("Пагинация", "LIMIT 10 и OFFSET", "120 страниц для 1200 товаров"),
                    ("Неправильный ввод", "Схемы и форма", "Статус 422; заполненные значения сохранены"),
                    ("Данные после старта", "Файловая SQLite", "Повторное открытие не удаляет записи"),
                ], (36, 54, 80))
            elif line.startswith("Для поиска используется"):
                table("Таблица 3 — Основные поля модели данных", ("Сущность и поле", "Представление", "Правило"), [
                    ("Category.id", "Целое", "Первичный ключ"),
                    ("Category.name", "Строка до 100 символов", "Непустое; уникальна нормализованная форма"),
                    ("Category.description", "Текст или NULL", "Не более 2000 символов"),
                    ("Product.id", "Целое", "Первичный ключ"),
                    ("Product.name", "Строка до 200 символов", "Непустое; нормализованная копия для поиска"),
                    ("Product.sku", "Строка до 64 символов", "Уникальный; латинские буквы, цифры, дефис, подчёркивание"),
                    ("Product.category_id", "Целое", "Внешний ключ; категория существует; RESTRICT"),
                    ("Product.price_kopecks", "Целые копейки", "От 0 до 2^63 − 1; при вводе рубли строкой"),
                    ("Product.quantity", "Целое количество", "От 0 до 2^63 − 1"),
                    ("Product.description", "Текст или NULL", "Не более 2000 символов"),
                    ("Product.created_at, updated_at", "Дата и время", "Хранятся в UTC; обновление при изменении"),
                ], (52, 48, 70))
                document.add_paragraph("Связь таблиц: categories.id → products.category_id, один ко многим. Удаление занятой категории запрещено. Нормализованные поля служебные и не входят в публичную схему API.")
            elif line.startswith("Компоненты организованы"):
                table("Таблица 4 — Ответственность компонентов", ("Компонент", "Назначение"), [
                    ("database.py", "Движок, внешний ключ, сессия на запрос"),
                    ("models.py / schemas.py", "Таблицы, ограничения и проверка входных данных"),
                    ("catalog.py", "Общие операции, транзакции, поиск и фильтры"),
                    ("routers/api.py", "JSON-контракт и статусы HTTP"),
                    ("routers/pages.py", "HTML-формы, сообщения и перенаправления"),
                    ("templates / static", "Представление, оформление и местное время"),
                    ("demo_data.py / seed_demo.py", "Воспроизводимые примеры, backup и загрузка"),
                ], (65, 105))
            elif line.startswith("На этом наборе в Chromium"):
                table("Таблица 5 — Фактические результаты проверки", ("Набор", "Количество", "Результат"), [
                    ("База данных", "35 тестов", "Пройдены"),
                    ("JSON API", "81 тест", "Пройдены"),
                    ("HTML-страницы", "31 тест", "Пройдены"),
                    ("Загрузчик примеров", "7 тестов", "Пройдены"),
                    ("Браузер, основные сценарии", "14 сценариев", "Пройдены в Chromium"),
                    ("Браузер, 1200 товаров", "6 сценариев", "Пройдены в Chromium"),
                    ("SQLite большого набора", "integrity_check, foreign_key_check", "ok; ошибок связей нет"),
                ], (67, 45, 58))

    heading("Список использованных источников")
    sources = json.loads((DOCS / "coursework" / "bibliography.json").read_text())["sources"]
    for source in sources:
        paragraph = document.add_paragraph(f"{source['id']}. {source['description']}")
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.paragraph_format.space_after = WordPt(6)
    heading("Приложение А. Экраны и результаты демонстрации")
    images = [
        ("large-catalog/04-catalog-viewport.png", "А.1", "Каталог из 1200 демонстрационных товаров"),
        ("03-product-form.png", "А.2", "Форма редактирования товара"),
        ("04-validation.png", "А.3", "Сохранение ввода при неправильной цене"),
        ("07-category-protection.png", "А.4", "Запрет удаления категории с товарами"),
    ]
    for index, (filename, number, caption) in enumerate(images):
        if index: document.add_page_break()
        paragraph = document.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.add_run().add_picture(str(DOCS / "figures" / filename), width=Mm(150))
        paragraph = document.add_paragraph(f"Рисунок {number} — {caption}")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Mm(0)
        document.add_paragraph("Источник: собственная разработка; реальная браузерная проверка на демонстрационных данных.")
    # Ссылки на закладки сохраняют номера страниц при обновлении полей Word.
    for title, bookmark, level in headings:
        paragraph = toc_anchor.insert_paragraph_before()
        paragraph.paragraph_format.first_line_indent = Mm(0)
        paragraph.paragraph_format.left_indent = Mm(5) if level == 2 else Mm(0)
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(160), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        paragraph.add_run(title + "\t")
        field(paragraph, f" PAGEREF {bookmark} \\h ", "?")
    update = OxmlElement("w:updateFields"); update.set(qn("w:val"), "true")
    document.settings.element.append(update)
    document.core_properties.title = STUDENT["topic"]
    document.core_properties.author = STUDENT["student_name"]
    document.core_properties.subject = "Курсовая работа; Разработка программных модулей; 09.02.07"
    document.save(WORD_FILE)
    return WORD_FILE


def make_slides():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.core_properties.title = "Модуль управления каталогом товаров"
    deck.core_properties.author = STUDENT["student_name"]

    def box(slide, x, y, w, h, color=WHITE, line=None, rounded=True):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                      Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(color)
        if line: shape.line.color.rgb = RGBColor.from_string(line)
        else: shape.line.fill.background()
        return shape

    def text(slide, value, x, y, w, h, size=22, color=INK, bold=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = shape.text_frame; frame.word_wrap = True
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(value.split("\n")):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Arial"; paragraph.font.size = Pt(size)
            paragraph.font.bold = bold; paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = Pt(12)
        return shape

    def slide(title, subtitle, note):
        item = deck.slides.add_slide(deck.slide_layouts[6])
        item.background.fill.solid(); item.background.fill.fore_color.rgb = RGBColor.from_string(BACK)
        box(item, .45, .5, .13, .47, BLUE, rounded=False)
        text(item, title, .8, .48, 11.7, .7, 30, bold=True)
        text(item, subtitle, .8, 1.25, 11.7, .65, 16, GRAY)
        text(item, "МФЮА · Учебный проект · 2026", .8, 7.05, 10, .25, 11, GRAY)
        text(item, f"{len(deck.slides):02d} / 08", 11.7, 7.05, .8, .25, 11, GRAY)
        item.notes_slide.notes_text_frame.text = note
        return item

    s = slide("Модуль управления каталогом товаров", "Курсовая работа · Python / FastAPI / SQLite", "Представиться, назвать тему и цель. Автор: Чемоданов Никита Сергеевич, группа ИПо-41. Это действующее локальное приложение, а не внедрённая промышленная система.")
    text(s, "Товары, категории,\nцены и остатки", .85, 2.3, 7.5, 2.1, 44, bold=True)
    text(s, f"{STUDENT['student_name']} · {STUDENT['group_short']}\n{STUDENT['discipline']}\nРуководитель: {STUDENT['supervisor']} · {STUDENT['city']}, {STUDENT['year']}", .85, 4.85, 11.6, 1.75, 17, GRAY)
    box(s, 9.4, 2.45, 2.9, 2.9, BLUE)
    text(s, "К", 10.22, 2.7, 1.3, 1.6, 92, WHITE, True)

    s = slide("Цель и задачи", "Менеджер ведёт ассортимент через браузер", "Пояснить границы: товары и категории, без оплаты и доставки. Поиск и фильтры можно комбинировать; остаток — текущее число, не журнал склада.")
    for i, (title, body) in enumerate([
        ("01  Данные", "Создавать и изменять товары\nи категории, просматривать карточки."),
        ("02  Поиск", "Отбирать по названию, категории,\nцене и наличию; сортировать."),
        ("03  Корректность", "Проверять ввод, уникальность\nи связи; подтверждать удаление."),
    ]):
        x = .8 + i * 4.13
        box(s, x, 2.25, 3.9, 3.35)
        text(s, title, x+.25, 2.65, 3.4, .6, 24, BLUE, True)
        text(s, body, x+.25, 3.55, 3.4, 1.55, 21)

    s = slide("Модель данных", "Одна категория — много товаров", "Объяснить идентификаторы и внешний ключ. Цена в копейках исключает float, количество неотрицательно. Категория с товарами не удаляется. Иерархия категорий не реализована.")
    box(s, .85, 2.2, 4.5, 3.9)
    text(s, "КАТЕГОРИЯ", 1.15, 2.55, 3.9, .5, 24, BLUE, True)
    text(s, "id — первичный ключ\nНазвание — уникальное\nОписание", 1.15, 3.35, 3.9, 2.0, 22)
    box(s, 7.35, 2.2, 5.1, 3.9)
    text(s, "ТОВАР", 7.65, 2.55, 4.4, .5, 24, BLUE, True)
    text(s, "id · category_id · название\nУникальный артикул\nЦена в копейках · остаток\nДаты создания и изменения", 7.65, 3.35, 4.4, 2.2, 21)
    line = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(5.35), Inches(4.1), Inches(7.35), Inches(4.1))
    line.line.color.rgb = RGBColor.from_string(BLUE); line.line.width = Pt(2)
    text(s, "1 : N", 5.8, 3.5, 1.2, .4, 20, BLUE, True)

    s = slide("Как устроено приложение", "HTML и JSON API используют общие правила", "Проследить запрос: страница или API вызывает catalog.py. SQLAlchemy работает с SQLite. После успешного POST выполняется 303 и GET. При конфликте транзакция откатывается.")
    nodes = [("Браузер", "Формы и страницы"), ("FastAPI", "HTML / JSON"), ("catalog.py", "Проверки и операции"), ("SQLite", "SQLAlchemy / Session")]
    for i, (title, caption) in enumerate(nodes):
        x=.8+i*3.15
        box(s,x,2.65,2.85,2.25,BLUE if i==2 else WHITE)
        text(s,title,x+.2,3.0,2.45,.55,24,WHITE if i==2 else BLUE,True)
        text(s,caption,x+.2,3.85,2.45,.7,17,WHITE if i==2 else GRAY)
        if i<3:text(s,"→",x+2.91,3.4,.23,.5,24,GRAY)
    text(s,"Входные схемы → ограничения моделей → ограничения базы",.85,5.6,11.8,.8,23,bold=True)

    s = slide("Веб-интерфейс", "Поиск, фильтры и пагинация; адаптивные формы", "Показать реальный список. Фильтры применяются вместе, сохраняются в URL и ссылках страниц. Подчеркнуть, что интерфейс использует локальные стили и основные формы работают без JavaScript.")
    text(s,"• Товары и категории\n• Поиск и фильтры\n• Карточки и формы\n• Подтверждение удаления",.85,2.5,3.15,3.4,23)
    s.shapes.add_picture(str(DOCS/'figures/large-catalog/05-catalog-slide.png'), Inches(4.35), Inches(2.0), width=Inches(8.15))

    s = slide("Демонстрационный ассортимент", "Синтетические данные для функциональной демонстрации", "1200 товаров в 12 категориях. Цены и остатки условные. Загрузчик сохраняет backup, пропускает существующие артикулы, добавляет недостающее одной транзакцией. Повторный запуск не возвращает изменённую цену к исходной.")
    for i,(number,label) in enumerate([("1200","товаров"),("12","категорий"),("133","без наличия"),("120","страниц")]):
        x=.8+i*3.15
        box(s,x,2.2,2.85,2.2)
        text(s,number,x+.25,2.55,2.35,.85,44,BLUE,True)
        text(s,label,x+.25,3.55,2.35,.45,19,GRAY)
    text(s,"Резервная копия → одна транзакция → без перезаписи записей",.85,5.05,11.6,.9,23,bold=True)
    text(s,"Команда: python -m scripts.seed_demo",.85,6.0,11.6,.5,20,GRAY)

    s = slide("Проверка результата", "154 автоматических теста; реальные браузерные сценарии", "Проверки базы, API, HTML и загрузчика. Отдельно 14 сценариев интерфейса и 6 на большом наборе. Есть одно известное предупреждение TestClient/HTTPX. Число проверок не означает отсутствие всех ошибок и не заменяет нагрузочное тестирование.")
    for i,(number,label) in enumerate([("35","База данных"),("81","JSON API"),("31","HTML"),("7","Загрузчик")]):
        x=.8+i*3.15
        box(s,x,2.2,2.85,2.05)
        text(s,number,x+.25,2.5,2.35,.85,42,BLUE,True)
        text(s,label,x+.25,3.45,2.35,.45,18,GRAY)
    text(s,"14 сценариев интерфейса + 6 проверок большого каталога",.85,4.85,11.7,.8,25,bold=True)
    text(s,"Исправлены переполнение мобильной страницы и выравнивание таблицы.\nОтдельно проверены повторная загрузка, откат и резервная копия WAL.",.85,5.8,11.7,.9,19,GRAY)

    s = slide("Результат и развитие", "Согласованные функции каталога реализованы", "Завершить перечислением результата. Обозначить отсутствие авторизации, миграций и журнала остатков. Дальнейшее развитие — роли пользователей, история операций и проверка одновременной работы. Приступить к короткой живой демонстрации.")
    box(s,.8,2.2,5.95,3.85)
    text(s,"РЕАЛИЗОВАНО",1.1,2.55,5.3,.5,23,BLUE,True)
    text(s,"Товары и категории\nКарточки, поиск, фильтры\nПроверка ввода и целостности\nВоспроизводимые тесты и данные",1.1,3.35,5.3,2.3,23)
    box(s,7.05,2.2,5.45,3.85)
    text(s,"НАПРАВЛЕНИЯ РАЗВИТИЯ",7.35,2.55,4.85,.5,21,BLUE,True)
    text(s,"Роли пользователей\nИстория движения остатков\nМиграции схемы\nМногопользовательские проверки",7.35,3.35,4.85,2.3,23)
    SLIDES_FILE.parent.mkdir(parents=True,exist_ok=True)
    deck.save(SLIDES_FILE)
    return SLIDES_FILE


if __name__ == "__main__":
    print(make_word())
    print(make_slides())
